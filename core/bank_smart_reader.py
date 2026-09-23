from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from typing import Iterable

from openpyxl import load_workbook

from config.column_mapping import BANK_HISTORY_HEADERS, SALES_SMART_HEADERS
from core.models import (
    BankHistoryData,
    BankTransactionRecord,
    ExcelFile,
    ReportPeriod,
    SalesInvoiceRecord,
    SalesSmartData,
    Severity,
    ValidationIssue,
)
from core.normalization import clean_text, fold_text, money_to_int, parse_date


def _load_read_only(source: ExcelFile):
    if not source.name.lower().endswith(".xlsx"):
        raise ValueError(f"{source.name}: phiên bản đầu chỉ hỗ trợ file .xlsx")
    try:
        return load_workbook(BytesIO(source.content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Không thể mở file {source.name}: {exc}") from exc


def _aliases(logical_name: str, mapping: dict[str, tuple[str, ...]]) -> set[str]:
    return {fold_text(alias) for alias in mapping[logical_name]}


def _matching_columns(headers: list[object], aliases: Iterable[str]) -> list[int]:
    folded_aliases = {fold_text(alias) for alias in aliases}
    exact = [
        index
        for index, value in enumerate(headers)
        if fold_text(value) in folded_aliases
    ]
    if exact:
        return exact
    return [
        index
        for index, value in enumerate(headers)
        if any(alias in fold_text(value) for alias in folded_aliases)
    ]


def _first_column(headers: list[object], mapping: dict[str, tuple[str, ...]], logical_name: str):
    columns = _matching_columns(headers, mapping[logical_name])
    return columns[0] if columns else None


def _header_candidate(workbook, required_logical_names, mapping, max_rows: int = 30):
    for worksheet in workbook.worksheets:
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, max_rows), values_only=True),
            start=1,
        ):
            headers = list(values)
            if all(_first_column(headers, mapping, logical_name) is not None for logical_name in required_logical_names):
                return worksheet, row_number, headers
    return None


def _data_rows(worksheet, header_row: int, blank_limit: int = 100):
    """Dừng sau vùng dữ liệu thật, kể cả khi Excel có định dạng dư hàng trăm nghìn dòng."""
    consecutive_blank = 0
    saw_data = False
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        row = tuple(values)
        if not any(value not in (None, "") for value in row):
            consecutive_blank += 1
            if saw_data and consecutive_blank >= blank_limit:
                break
            continue
        saw_data = True
        consecutive_blank = 0
        yield row_number, row


def _value(row: tuple, column: int | None):
    if column is None or column >= len(row):
        return None
    return row[column]


def _issue(severity, code, message, source, sheet="", row=None):
    return ValidationIssue(severity, code, message, source.name, sheet, row)


def _direction_marker(value: object) -> str:
    text = clean_text(value).replace("'", "").strip()
    return text[:1] if text[:1] in {"+", "-"} else ""


def _unsigned_money(value: object) -> int:
    text = clean_text(value).replace("'", "").strip()
    if text[:1] in {"+", "-"}:
        text = text[1:]
    return money_to_int(text)


def _money_decimal(value: object) -> Decimal:
    if isinstance(value, bool) or value in (None, ""):
        raise ValueError("Số tiền không hợp lệ")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = clean_text(value).replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", text):
        text = text.replace(",", "")
    elif "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Không đọc được số tiền: {value}") from exc


def read_bank_history_file(source: ExcelFile) -> BankHistoryData:
    result = BankHistoryData(source_file=source.name)
    try:
        workbook = _load_read_only(source)
    except ValueError as exc:
        result.issues.append(_issue(Severity.ERROR, "INVALID_BANK_FILE", str(exc), source))
        return result

    candidate = _header_candidate(
        workbook,
        ("transaction_datetime", "amount", "content"),
        BANK_HISTORY_HEADERS,
    )
    if candidate is None:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "BANK_HEADER_NOT_FOUND",
                "Không tìm thấy tiêu đề ngày giao dịch, số tiền giao dịch và nội dung",
                source,
            )
        )
        return result

    worksheet, header_row, headers = candidate
    date_column = _first_column(headers, BANK_HISTORY_HEADERS, "transaction_datetime")
    content_column = _first_column(headers, BANK_HISTORY_HEADERS, "content")
    transaction_id_column = _first_column(headers, BANK_HISTORY_HEADERS, "transaction_id")
    counterparty_column = _first_column(headers, BANK_HISTORY_HEADERS, "counterparty_account")
    currency_column = _first_column(headers, BANK_HISTORY_HEADERS, "currency")
    credit_column = _first_column(headers, BANK_HISTORY_HEADERS, "credit")
    amount_columns = _matching_columns(headers, BANK_HISTORY_HEADERS["amount"])

    rows = list(_data_rows(worksheet, header_row))
    signed_scores = {
        column: sum(_direction_marker(_value(row, column)) != "" for _, row in rows[:200])
        for column in amount_columns
    }
    signed_column = max(signed_scores, key=signed_scores.get) if signed_scores else None
    if signed_column is not None and signed_scores[signed_column] == 0:
        signed_column = None
    numeric_candidates = [column for column in amount_columns if column != signed_column]
    amount_column = credit_column if credit_column is not None else (numeric_candidates[0] if numeric_candidates else signed_column)

    if amount_column is None:
        result.issues.append(
            _issue(Severity.ERROR, "BANK_AMOUNT_COLUMN_NOT_FOUND", "Không xác định được cột tiền vào", source, worksheet.title, header_row)
        )
        return result
    if signed_column is None and credit_column is None:
        result.issues.append(
            _issue(
                Severity.WARNING,
                "BANK_DIRECTION_COLUMN_NOT_FOUND",
                "Không có cột dấu +/- hoặc Ghi có; tool coi các số tiền dương là tiền vào",
                source,
                worksheet.title,
                header_row,
            )
        )

    seen_keys = set()
    dates: list[date] = []
    for row_number, row in rows:
        if signed_column is not None and _direction_marker(_value(row, signed_column)) != "+":
            continue
        currency = fold_text(_value(row, currency_column))
        if currency and currency not in {"VND", "VN D"}:
            result.issues.append(
                _issue(
                    Severity.WARNING,
                    "BANK_NON_VND_SKIPPED",
                    f"Bỏ qua giao dịch không phải VND tại dòng {row_number}",
                    source,
                    worksheet.title,
                    row_number,
                )
            )
            continue
        transaction_date = parse_date(_value(row, date_column))
        if transaction_date is None:
            result.issues.append(
                _issue(Severity.ERROR, "BANK_INVALID_DATE", "Không đọc được ngày giao dịch", source, worksheet.title, row_number)
            )
            continue
        try:
            amount = _unsigned_money(_value(row, amount_column))
        except ValueError as exc:
            result.issues.append(
                _issue(Severity.ERROR, "BANK_INVALID_AMOUNT", str(exc), source, worksheet.title, row_number)
            )
            continue
        if amount <= 0:
            continue
        content = clean_text(_value(row, content_column))
        transaction_id = clean_text(_value(row, transaction_id_column))
        key = transaction_id or (transaction_date, amount, content)
        if key in seen_keys:
            result.issues.append(
                _issue(Severity.WARNING, "BANK_DUPLICATE_TRANSACTION", "Giao dịch ngân hàng trùng đã được loại bỏ", source, worksheet.title, row_number)
            )
            continue
        seen_keys.add(key)
        if not content:
            result.issues.append(
                _issue(Severity.WARNING, "BANK_EMPTY_CONTENT", "Giao dịch tiền vào không có nội dung chuyển khoản", source, worksheet.title, row_number)
            )
        result.transactions.append(
            BankTransactionRecord(
                transaction_date=transaction_date,
                amount=amount,
                content=content,
                transaction_id=transaction_id,
                counterparty_account=clean_text(_value(row, counterparty_column)),
                source_file=source.name,
                sheet=worksheet.title,
                source_row=row_number,
            )
        )
        dates.append(transaction_date)

    if dates:
        result.period = ReportPeriod(min(dates), max(dates))
    else:
        result.issues.append(
            _issue(Severity.ERROR, "BANK_NO_INCOMING", "Không tìm thấy giao dịch tiền vào hợp lệ", source, worksheet.title)
        )
    return result


def read_sales_smart_file(source: ExcelFile) -> SalesSmartData:
    result = SalesSmartData(source_file=source.name)
    try:
        workbook = _load_read_only(source)
    except ValueError as exc:
        result.issues.append(_issue(Severity.ERROR, "INVALID_SALES_SMART_FILE", str(exc), source))
        return result

    candidate = _header_candidate(
        workbook,
        ("document_id", "invoice_date", "total_amount"),
        SALES_SMART_HEADERS,
    )
    if candidate is None:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "SALES_SMART_HEADER_NOT_FOUND",
                "Không tìm thấy tiêu đề ID chứng từ, ngày HĐ và Tổng Tiền",
                source,
            )
        )
        return result

    worksheet, header_row, headers = candidate
    columns = {
        logical_name: _first_column(headers, SALES_SMART_HEADERS, logical_name)
        for logical_name in SALES_SMART_HEADERS
    }
    groups = defaultdict(
        lambda: {
            "invoice_number": "",
            "invoice_date": None,
            "customer_name": "",
            "debit_accounts": set(),
            "amount": Decimal("0"),
            "rows": [],
        }
    )

    for row_number, row in _data_rows(worksheet, header_row):
        document_type = fold_text(_value(row, columns["document_type"]))
        if document_type and document_type != "HDBR":
            continue
        document_id = clean_text(_value(row, columns["document_id"]))
        if not document_id:
            continue
        invoice_date = parse_date(_value(row, columns["invoice_date"])) or parse_date(
            _value(row, columns["document_date"])
        )
        if invoice_date is None:
            result.issues.append(
                _issue(Severity.ERROR, "SALES_SMART_INVALID_DATE", f"Chứng từ {document_id} không đọc được ngày hóa đơn", source, worksheet.title, row_number)
            )
            continue
        group = groups[document_id]
        invoice_number = clean_text(_value(row, columns["invoice_number"]))
        customer_name = clean_text(_value(row, columns["customer_name"]))
        if group["invoice_date"] and group["invoice_date"] != invoice_date:
            result.issues.append(
                _issue(Severity.WARNING, "SALES_SMART_INCONSISTENT_DATE", f"Chứng từ {document_id} có nhiều ngày hóa đơn", source, worksheet.title, row_number)
            )
        if group["customer_name"] and customer_name and group["customer_name"] != customer_name:
            result.issues.append(
                _issue(Severity.WARNING, "SALES_SMART_INCONSISTENT_CUSTOMER", f"Chứng từ {document_id} có nhiều tên khách hàng", source, worksheet.title, row_number)
            )
        group["invoice_date"] = group["invoice_date"] or invoice_date
        group["invoice_number"] = group["invoice_number"] or invoice_number
        group["customer_name"] = group["customer_name"] or customer_name
        debit_account = clean_text(_value(row, columns["debit_account"]))
        if debit_account:
            group["debit_accounts"].add(debit_account)
        raw_amount = _value(row, columns["total_amount"])
        if raw_amount not in (None, ""):
            try:
                group["amount"] += _money_decimal(raw_amount)
            except ValueError as exc:
                result.issues.append(
                    _issue(Severity.ERROR, "SALES_SMART_INVALID_AMOUNT", f"Chứng từ {document_id}: {exc}", source, worksheet.title, row_number)
                )
        group["rows"].append(row_number)

    dates = []
    for document_id, group in groups.items():
        invoice_date = group["invoice_date"]
        if invoice_date is None:
            continue
        raw_total = group["amount"]
        amount = int(raw_total.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        if raw_total != Decimal(amount):
            result.issues.append(
                _issue(
                    Severity.WARNING,
                    "SALES_SMART_TOTAL_ROUNDED",
                    f"Chứng từ {document_id} có tổng tiền lẻ và đã làm tròn về VND",
                    source,
                    worksheet.title,
                    group["rows"][0],
                )
            )
        if amount == 0:
            result.issues.append(
                _issue(Severity.WARNING, "SALES_SMART_ZERO_TOTAL", f"Chứng từ {document_id} có tổng tiền bằng 0", source, worksheet.title, group["rows"][0])
            )
        result.invoices.append(
            SalesInvoiceRecord(
                document_id=document_id,
                invoice_number=group["invoice_number"],
                invoice_date=invoice_date,
                amount=amount,
                customer_name=group["customer_name"],
                debit_account="; ".join(sorted(group["debit_accounts"])),
                source_file=source.name,
                sheet=worksheet.title,
                source_rows=tuple(group["rows"]),
            )
        )
        dates.append(invoice_date)

    result.invoices.sort(key=lambda record: (record.invoice_date, record.invoice_number, record.document_id))
    if dates:
        result.period = ReportPeriod(min(dates), max(dates))
    else:
        result.issues.append(
            _issue(Severity.ERROR, "SALES_SMART_NO_INVOICES", "Không tìm thấy hóa đơn bán ra hợp lệ", source, worksheet.title)
        )
    return result

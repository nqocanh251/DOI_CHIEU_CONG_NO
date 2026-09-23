from __future__ import annotations

from collections import Counter
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.models import BankHistoryData, BankSmartReconciliation, SalesSmartData
from core.normalization import format_date
from core.report_exporter import (
    BLUE,
    BLUE_LIGHT,
    GREEN,
    GREEN_TEXT,
    GRID,
    MONEY_FORMAT,
    NAVY,
    REPORT_BODY_FONT_SIZE,
    REPORT_SECTION_FONT_SIZE,
    REPORT_TITLE_FONT_SIZE,
    RED,
    RED_TEXT,
    SLATE_LIGHT,
    WHITE,
    YELLOW,
    YELLOW_TEXT,
    _append_dataframe,
    _set_report_printing,
    _style_sheet,
    issues_to_dataframe,
)


def bank_smart_results_to_dataframe(reconciliation: BankSmartReconciliation) -> pd.DataFrame:
    columns = [
        "STT", "Trạng thái", "Ngày ngân hàng", "ID giao dịch ngân hàng",
        "Số tiền ngân hàng", "Nội dung chuyển khoản", "Tài khoản đối ứng",
        "Ngày Smart", "ID chứng từ Smart", "Số hóa đơn", "Tên khách hàng Smart",
        "TK Nợ Smart", "Tổng tiền Smart", "Chênh lệch", "Số ngày chênh lệch",
        "Số ứng viên Smart", "Số HĐ ứng viên", "Khách hàng ứng viên",
        "Nguyên nhân", "File ngân hàng", "Dòng ngân hàng", "File Smart", "Dòng Smart",
    ]
    rows = []
    for index, result in enumerate(reconciliation.results, 1):
        bank = result.bank
        smart = result.smart
        rows.append(
            {
                "STT": index,
                "Trạng thái": result.status,
                "Ngày ngân hàng": format_date(bank.transaction_date),
                "ID giao dịch ngân hàng": bank.transaction_id,
                "Số tiền ngân hàng": bank.amount,
                "Nội dung chuyển khoản": bank.content,
                "Tài khoản đối ứng": bank.counterparty_account,
                "Ngày Smart": format_date(smart.invoice_date) if smart else "",
                "ID chứng từ Smart": smart.document_id if smart else "",
                "Số hóa đơn": smart.invoice_number if smart else "",
                "Tên khách hàng Smart": smart.customer_name if smart else "",
                "TK Nợ Smart": smart.debit_account if smart else "",
                "Tổng tiền Smart": smart.amount if smart else None,
                "Chênh lệch": result.amount_difference,
                "Số ngày chênh lệch": result.day_difference,
                "Số ứng viên Smart": result.candidate_count,
                "Số HĐ ứng viên": "; ".join(result.candidate_invoice_numbers),
                "Khách hàng ứng viên": "; ".join(result.candidate_customer_names),
                "Nguyên nhân": result.reason,
                "File ngân hàng": bank.source_file,
                "Dòng ngân hàng": bank.source_row,
                "File Smart": smart.source_file if smart else "",
                "Dòng Smart": ", ".join(map(str, smart.source_rows)) if smart else "",
            }
        )
    dataframe = pd.DataFrame(rows, columns=columns)
    dataframe["Số ngày chênh lệch"] = pd.array(
        dataframe["Số ngày chênh lệch"], dtype="Int64"
    )
    return dataframe


def _summary_value_style(cell, *, money: bool = False, status: bool = False):
    cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color="263746")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if money:
        cell.number_format = MONEY_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center")
    if status:
        normalized = str(cell.value or "").upper()
        if normalized == "KHỚP HOÀN TOÀN":
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=GREEN_TEXT, bold=True)
        elif normalized == "CẦN KIỂM TRA":
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=YELLOW_TEXT, bold=True)
        else:
            cell.fill = PatternFill("solid", fgColor=RED)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=RED_TEXT, bold=True)


def _add_summary_section(ws, row: int, title: str, items: list[tuple[str, object, str]]) -> int:
    thin_grid = Side(style="thin", color=GRID)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    section = ws.cell(row, 1, title)
    section.fill = PatternFill("solid", fgColor=BLUE)
    section.font = Font(name="Segoe UI", size=REPORT_SECTION_FONT_SIZE, color=WHITE, bold=True)
    section.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 26
    row += 1
    for offset, (label, value, kind) in enumerate(items):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        label_cell = ws.cell(row, 1, label)
        value_cell = ws.cell(row, 2, value)
        label_cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
        label_cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=NAVY, bold=True)
        label_cell.alignment = Alignment(vertical="center", wrap_text=True)
        value_cell.fill = PatternFill("solid", fgColor=SLATE_LIGHT if offset % 2 else WHITE)
        _summary_value_style(value_cell, money=kind == "money", status=kind == "status")
        for column in range(1, 7):
            ws.cell(row, column).border = Border(bottom=thin_grid)
        ws.row_dimensions[row].height = 32
        row += 1
    return row + 1


def _build_summary(ws, bank: BankHistoryData, smart: SalesSmartData, reconciliation: BankSmartReconciliation):
    ws.merge_cells("A1:F1")
    ws["A1"] = "BÁO CÁO ĐỐI CHIẾU NGÂN HÀNG ↔ SMART"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Segoe UI", size=REPORT_TITLE_FONT_SIZE, color=WHITE, bold=True)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        f"Chỉ đối chiếu ngày và số tiền · Ngày Smart trễ tối đa 1 ngày · Chênh lệch tối đa 1 đồng · "
        f"Xuất lúc {datetime.now():%d/%m/%Y %H:%M:%S}"
    )
    ws["A2"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    ws["A2"].font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=NAVY, bold=True)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    statuses = Counter(result.status for result in reconciliation.results)
    bank_total = sum(record.amount for record in bank.transactions)
    paired_results = [result for result in reconciliation.results if result.smart is not None]
    smart_total = sum(result.smart.amount for result in paired_results)
    result_status = (
        "KHỚP HOÀN TOÀN"
        if reconciliation.results and all(result.status == "KHỚP" for result in reconciliation.results)
        else "CẦN KIỂM TRA"
    )

    row = 4
    row = _add_summary_section(
        ws,
        row,
        "THÔNG TIN ĐỐI CHIẾU",
        [
            ("File ngân hàng", bank.source_file, "text"),
            ("Kỳ ngân hàng", bank.period.label, "text"),
            ("File Smart", smart.source_file, "text"),
            ("Kỳ Smart", smart.period.label, "text"),
            ("Phạm vi Smart", "Toàn bộ chứng từ HDBR; TK Nợ chỉ hiển thị tham khảo", "text"),
        ],
    )
    row = _add_summary_section(
        ws,
        row,
        "TỔNG HỢP SỐ LIỆU",
        [
            ("Số giao dịch tiền vào ngân hàng", len(bank.transactions), "count"),
            ("Tổng tiền vào ngân hàng", bank_total, "money"),
            ("Số hóa đơn Smart", len(smart.invoices), "count"),
            ("Tổng tiền Smart đã ghép duy nhất", smart_total, "money"),
            ("Kết quả tổng thể", result_status, "status"),
        ],
    )
    status_items = [(status, count, "count") for status, count in sorted(statuses.items())]
    row = _add_summary_section(ws, row, "SỐ LƯỢNG THEO TRẠNG THÁI", status_items)

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22
    for column in "CDEF":
        ws.column_dimensions[column].width = 14
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.print_area = f"A1:F{row - 1}"
    _set_report_printing(ws, landscape=False)


def build_bank_smart_excel_report(
    bank: BankHistoryData,
    smart: SalesSmartData,
    reconciliation: BankSmartReconciliation,
) -> bytes:
    workbook = Workbook()
    workbook.properties.title = "Báo cáo đối chiếu ngân hàng – Smart"
    workbook.properties.subject = "Đối chiếu giao dịch tiền vào với hóa đơn bán ra Smart"
    workbook.properties.creator = "Tool đối chiếu công nợ"

    summary = workbook.active
    summary.title = "Tong_quan"
    _build_summary(summary, bank, smart, reconciliation)

    detail = workbook.create_sheet("Ngan_hang_Smart")
    _append_dataframe(detail, bank_smart_results_to_dataframe(reconciliation))
    _style_sheet(detail)
    detail.sheet_properties.tabColor = GREEN_TEXT

    issues = workbook.create_sheet("Loi_du_lieu")
    _append_dataframe(issues, issues_to_dataframe(reconciliation.issues))
    _style_sheet(issues, status_column_name="Mức độ")
    issues.sheet_properties.tabColor = RED_TEXT

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

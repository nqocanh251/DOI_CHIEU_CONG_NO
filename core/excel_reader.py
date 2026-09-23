from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from io import BytesIO
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from config.column_mapping import HDDT_REQUIRED_HEADERS, KIOT_REQUIRED_HEADERS
from core.models import (
    ExcelFile,
    HddtData,
    InvoiceRecord,
    KiotData,
    PaymentRecord,
    ReportPeriod,
    Severity,
    SmartData,
    ValidationIssue,
)
from core.normalization import (
    clean_text,
    fold_text,
    money_to_int,
    normalize_invoice_number,
    normalize_status,
    normalize_symbol,
    normalize_tax_id,
    parse_date,
    parse_report_period,
    parse_time,
)


class WorkbookFormatError(ValueError):
    pass


def _load(source: ExcelFile):
    if not source.name.lower().endswith(".xlsx"):
        raise WorkbookFormatError(f"{source.name}: phiên bản đầu chỉ hỗ trợ file .xlsx")
    try:
        return load_workbook(BytesIO(source.content), data_only=True, read_only=False)
    except Exception as exc:
        raise WorkbookFormatError(f"Không thể mở file {source.name}: {exc}") from exc


def _find_period(workbook) -> ReportPeriod:
    for ws in workbook.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
            for value in row:
                start, end = parse_report_period(value)
                if start and end:
                    return ReportPeriod(start, end)
    return ReportPeriod()


def _find_header_row(ws: Worksheet, required_terms: Iterable[str], max_scan_rows: int = 30) -> int | None:
    required = [fold_text(term) for term in required_terms]
    for row_number in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = [fold_text(ws.cell(row_number, column).value) for column in range(1, ws.max_column + 1)]
        if all(any(term in value for value in values) for term in required):
            return row_number
    return None


def _header_map(ws: Worksheet, header_row: int, aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    values = {column: fold_text(ws.cell(header_row, column).value) for column in range(1, ws.max_column + 1)}
    result: dict[str, int] = {}
    for logical_name, candidates in aliases.items():
        normalized = [fold_text(candidate) for candidate in candidates]
        for column, value in values.items():
            if any(candidate == value or candidate in value for candidate in normalized):
                result[logical_name] = column
                break
    return result


def _issue(
    severity: Severity,
    code: str,
    message: str,
    source: ExcelFile,
    sheet: str = "",
    row: int | None = None,
) -> ValidationIssue:
    return ValidationIssue(severity, code, message, source.name, sheet, row)


def read_hddt_files(sources: list[ExcelFile]) -> HddtData:
    result = HddtData(files=[source.name for source in sources])
    if not sources:
        result.issues.append(
            ValidationIssue(Severity.ERROR, "NO_HDDT_FILE", "Chưa chọn file HĐĐT")
        )
        return result

    records: list[InvoiceRecord] = []
    for source in sources:
        try:
            workbook = _load(source)
        except WorkbookFormatError as exc:
            result.issues.append(_issue(Severity.ERROR, "INVALID_HDDT_FILE", str(exc), source))
            continue

        period = _find_period(workbook)
        result.periods.append(period)
        candidate: tuple[Worksheet, int] | None = None
        for ws in workbook.worksheets:
            header_row = _find_header_row(ws, ["Số hóa đơn", "Tổng tiền thanh toán", "MST người bán"])
            if header_row:
                candidate = ws, header_row
                break
        if candidate is None:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "HDDT_HEADER_NOT_FOUND",
                    "Không tìm thấy dòng tiêu đề HĐĐT có Số hóa đơn, MST người bán và Tổng tiền thanh toán",
                    source,
                )
            )
            continue

        ws, header_row = candidate
        columns = _header_map(ws, header_row, HDDT_REQUIRED_HEADERS)
        missing = [name for name in HDDT_REQUIRED_HEADERS if name not in columns]
        if missing:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "HDDT_MISSING_COLUMNS",
                    "Thiếu cột bắt buộc: " + ", ".join(missing),
                    source,
                    ws.title,
                    header_row,
                )
            )
            continue

        parsed_count = 0
        for row_number in range(header_row + 1, ws.max_row + 1):
            raw_invoice = ws.cell(row_number, columns["invoice_number"]).value
            raw_tax_id = ws.cell(row_number, columns["supplier_tax_id"]).value
            if raw_invoice in (None, "") and raw_tax_id in (None, ""):
                continue
            invoice_number = normalize_invoice_number(raw_invoice)
            tax_id = normalize_tax_id(raw_tax_id)
            if not invoice_number or not tax_id:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "HDDT_INVALID_KEY",
                        "Dòng hóa đơn thiếu số hóa đơn hoặc MST người bán",
                        source,
                        ws.title,
                        row_number,
                    )
                )
                continue
            invoice_date = parse_date(ws.cell(row_number, columns["invoice_date"]).value)
            if invoice_date is None:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "HDDT_INVALID_DATE",
                        f"Không đọc được ngày của hóa đơn {invoice_number}",
                        source,
                        ws.title,
                        row_number,
                    )
                )
                continue
            try:
                amount = money_to_int(ws.cell(row_number, columns["amount"]).value)
            except ValueError as exc:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "HDDT_INVALID_AMOUNT",
                        f"Hóa đơn {invoice_number}: {exc}",
                        source,
                        ws.title,
                        row_number,
                    )
                )
                continue
            status = normalize_status(ws.cell(row_number, columns["status"]).value)
            if not status:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "HDDT_MISSING_STATUS",
                        f"Hóa đơn {invoice_number} thiếu trạng thái hóa đơn",
                        source,
                        ws.title,
                        row_number,
                    )
                )
                continue
            records.append(
                InvoiceRecord(
                    supplier_tax_id=tax_id,
                    supplier_name=clean_text(ws.cell(row_number, columns["supplier_name"]).value),
                    invoice_number=invoice_number,
                    invoice_number_raw=clean_text(raw_invoice),
                    invoice_symbol=normalize_symbol(ws.cell(row_number, columns["invoice_symbol"]).value),
                    invoice_date=invoice_date,
                    amount=amount,
                    status=status,
                    source_file=source.name,
                    sheet=ws.title,
                    source_row=row_number,
                )
            )
            parsed_count += 1
        if not parsed_count:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "HDDT_NO_DATA",
                    "Không tìm thấy dòng hóa đơn hợp lệ",
                    source,
                    ws.title,
                )
            )

    dedupe_key_to_record: dict[tuple, InvoiceRecord] = {}
    for record in records:
        key = (
            record.supplier_tax_id,
            record.invoice_symbol,
            record.invoice_number,
            record.invoice_date,
            record.amount,
            record.status,
        )
        existing = dedupe_key_to_record.get(key)
        if existing:
            result.issues.append(
                ValidationIssue(
                    Severity.WARNING,
                    "DUPLICATE_HDDT_ROW_REMOVED",
                    f"Đã loại bản ghi HĐĐT trùng hóa đơn {record.invoice_number}; bản đầu ở {existing.source_file}",
                    record.source_file,
                    record.sheet,
                    record.source_row,
                )
            )
        else:
            dedupe_key_to_record[key] = record
    result.invoices = list(dedupe_key_to_record.values())
    return result


def _smart_sheet(workbook) -> Worksheet | None:
    if "IN" in workbook.sheetnames:
        ws = workbook["IN"]
        if _find_header_row(ws, ["Chứng từ", "Hóa đơn"], 20):
            return ws
    for ws in workbook.worksheets:
        if _find_header_row(ws, ["Chứng từ", "Hóa đơn"], 20):
            return ws
    return None


def _smart_supplier(ws: Worksheet) -> tuple[str, str]:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
        for value in row:
            text = clean_text(value)
            if fold_text(text).startswith("DOI TUONG"):
                match = re.search(r"(\d{10,14})\s*-\s*(.+)$", text)
                if match:
                    return normalize_tax_id(match.group(1)), clean_text(match.group(2))
    return "", ""


def _combined_smart_headers(ws: Worksheet, top_row: int) -> dict[int, str]:
    headers: dict[int, str] = {}
    for column in range(1, ws.max_column + 1):
        parts = [clean_text(ws.cell(row, column).value) for row in range(top_row, min(top_row + 2, ws.max_row) + 1)]
        headers[column] = fold_text(" ".join(part for part in parts if part))
    return headers


def _find_column(headers: dict[int, str], predicate) -> int | None:
    for column, value in headers.items():
        if predicate(value):
            return column
    return None


def read_smart_file(source: ExcelFile) -> SmartData:
    result = SmartData(source_file=source.name)
    try:
        workbook = _load(source)
    except WorkbookFormatError as exc:
        result.issues.append(_issue(Severity.ERROR, "INVALID_SMART_FILE", str(exc), source))
        return result
    result.period = _find_period(workbook)
    ws = _smart_sheet(workbook)
    if ws is None:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "SMART_HEADER_NOT_FOUND",
                "Không tìm thấy sheet Smart có tiêu đề Chứng từ và Hóa đơn",
                source,
            )
        )
        return result

    result.supplier_tax_id, result.supplier_name = _smart_supplier(ws)
    if not result.supplier_tax_id:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "SMART_SUPPLIER_TAX_ID_NOT_FOUND",
                "Không đọc được MST tại dòng Đối tượng của Smart",
                source,
                ws.title,
            )
        )

    header_row = _find_header_row(ws, ["Chứng từ", "Hóa đơn"], 20)
    assert header_row is not None
    headers = _combined_smart_headers(ws, header_row)
    type_col = _find_column(headers, lambda v: "CHUNG TU" in v and "LOAI" in v)
    invoice_col = _find_column(headers, lambda v: "HOA DON" in v and " SO" in f" {v}")
    date_col = _find_column(headers, lambda v: "NGAY" in v)
    content_col = _find_column(headers, lambda v: "NOI DUNG" in v)
    amount_col = _find_column(headers, lambda v: "THANH TIEN" in v)
    paid_col = _find_column(headers, lambda v: "DA THANH TOAN" in v)
    if None in (type_col, invoice_col, date_col, content_col, amount_col, paid_col):
        result.issues.append(
            _issue(
                Severity.ERROR,
                "SMART_MISSING_COLUMNS",
                "Không xác định được đủ cột Loại, Số hóa đơn, Ngày, Nội dung, Thành tiền và Đã thanh toán",
                source,
                ws.title,
                header_row,
            )
        )
        return result

    type_col = int(type_col)
    invoice_col = int(invoice_col)
    date_col = int(date_col)
    content_col = int(content_col)
    amount_col = int(amount_col)
    paid_col = int(paid_col)
    reference_col = paid_col + 1 if paid_col + 1 <= ws.max_column else None

    current_invoice = ""
    current_invoice_raw = ""
    current_date: date | None = None
    current_block_row: int | None = None
    invoice_count = 0
    for row_number in range(header_row + 1, ws.max_row + 1):
        row_type = fold_text(ws.cell(row_number, type_col).value)
        raw_invoice = ws.cell(row_number, invoice_col).value
        invoice_number = normalize_invoice_number(raw_invoice)
        row_date = parse_date(ws.cell(row_number, date_col).value)
        content = clean_text(ws.cell(row_number, content_col).value)
        folded_content = fold_text(content)

        if row_type == "PNK" and invoice_number:
            if current_invoice and current_block_row is not None and invoice_number != current_invoice:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "SMART_INVOICE_TOTAL_NOT_FOUND",
                        f"Không tìm thấy dòng Cộng chứng từ cho hóa đơn {current_invoice}",
                        source,
                        ws.title,
                        current_block_row,
                    )
                )
            current_invoice = invoice_number
            current_invoice_raw = clean_text(raw_invoice)
            current_date = row_date
            current_block_row = row_number

        if folded_content.startswith("CONG CHUNG TU"):
            if not current_invoice:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "SMART_ORPHAN_TOTAL",
                        "Dòng Cộng chứng từ không liên kết được với số hóa đơn",
                        source,
                        ws.title,
                        row_number,
                    )
                )
            else:
                try:
                    amount = money_to_int(ws.cell(row_number, amount_col).value)
                except ValueError as exc:
                    result.issues.append(
                        _issue(
                            Severity.ERROR,
                            "SMART_INVALID_INVOICE_AMOUNT",
                            f"Hóa đơn {current_invoice}: {exc}",
                            source,
                            ws.title,
                            row_number,
                        )
                    )
                else:
                    result.invoices.append(
                        InvoiceRecord(
                            supplier_tax_id=result.supplier_tax_id,
                            supplier_name=result.supplier_name,
                            invoice_number=current_invoice,
                            invoice_number_raw=current_invoice_raw,
                            invoice_symbol="",
                            invoice_date=current_date,
                            amount=amount,
                            status="",
                            source_file=source.name,
                            sheet=ws.title,
                            source_row=row_number,
                            reference=content,
                        )
                    )
                    invoice_count += 1
            current_invoice = ""
            current_invoice_raw = ""
            current_date = None
            current_block_row = None

        paid_value = ws.cell(row_number, paid_col).value
        if row_date is not None and paid_value not in (None, "", 0, 0.0):
            try:
                paid_amount = money_to_int(paid_value)
            except ValueError as exc:
                result.issues.append(
                    _issue(
                        Severity.ERROR,
                        "SMART_INVALID_PAYMENT_AMOUNT",
                        str(exc),
                        source,
                        ws.title,
                        row_number,
                    )
                )
            else:
                if paid_amount > 0:
                    result.payments.append(
                        PaymentRecord(
                            source="SMART",
                            payment_date=row_date,
                            amount=paid_amount,
                            content=content,
                            source_file=source.name,
                            sheet=ws.title,
                            source_row=row_number,
                            reference=clean_text(ws.cell(row_number, reference_col).value) if reference_col else "",
                        )
                    )

    if current_invoice and current_block_row is not None:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "SMART_INVOICE_TOTAL_NOT_FOUND",
                f"Không tìm thấy dòng Cộng chứng từ cho hóa đơn {current_invoice}",
                source,
                ws.title,
                current_block_row,
            )
        )
    if not invoice_count:
        result.issues.append(
            _issue(Severity.ERROR, "SMART_NO_INVOICES", "Không đọc được hóa đơn nào trong Smart", source, ws.title)
        )
    if not result.payments:
        result.issues.append(
            _issue(
                Severity.WARNING,
                "SMART_NO_PAYMENTS",
                "Không tìm thấy dòng Đã thanh toán lớn hơn 0 trong Smart",
                source,
                ws.title,
            )
        )
    return result


def _kiot_sheet(workbook) -> tuple[Worksheet, int] | None:
    for ws in workbook.worksheets:
        header_row = _find_header_row(ws, ["Thời gian", "Diễn giải", "Ghi có"], 30)
        if header_row:
            return ws, header_row
    return None


def _kiot_supplier(ws: Worksheet, header_row: int) -> str:
    for row_number in range(1, header_row):
        for column in range(1, ws.max_column + 1):
            if fold_text(ws.cell(row_number, column).value) == "TEN NCC":
                for next_column in range(column + 1, min(ws.max_column, column + 3) + 1):
                    value = clean_text(ws.cell(row_number, next_column).value)
                    if value:
                        return value
    return ""


def read_kiot_file(source: ExcelFile, keyword: str) -> KiotData:
    result = KiotData(source_file=source.name, keyword=clean_text(keyword))
    if not result.keyword:
        result.issues.append(_issue(Severity.ERROR, "EMPTY_KEYWORD", "Từ khóa KIOT đang để trống", source))
        return result
    try:
        workbook = _load(source)
    except WorkbookFormatError as exc:
        result.issues.append(_issue(Severity.ERROR, "INVALID_KIOT_FILE", str(exc), source))
        return result
    result.period = _find_period(workbook)
    candidate = _kiot_sheet(workbook)
    if candidate is None:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "KIOT_HEADER_NOT_FOUND",
                "Không tìm thấy dòng tiêu đề KIOT có Thời gian, Diễn giải và Ghi có",
                source,
            )
        )
        return result
    ws, header_row = candidate
    result.supplier_name = _kiot_supplier(ws, header_row)
    columns = _header_map(ws, header_row, KIOT_REQUIRED_HEADERS)
    missing = [name for name in KIOT_REQUIRED_HEADERS if name not in columns]
    if missing:
        result.issues.append(
            _issue(
                Severity.ERROR,
                "KIOT_MISSING_COLUMNS",
                "Thiếu cột bắt buộc: " + ", ".join(missing),
                source,
                ws.title,
                header_row,
            )
        )
        return result

    folded_keyword = fold_text(result.keyword)
    transaction_for_keyword: dict[int, list[str]] = defaultdict(list)
    for row_number in range(header_row + 1, ws.max_row + 1):
        description = clean_text(ws.cell(row_number, columns["description"]).value)
        if folded_keyword not in fold_text(description):
            continue

        transaction_row: int | None = None
        for previous in range(row_number - 1, header_row, -1):
            date_value = parse_date(ws.cell(previous, columns["datetime"]).value)
            if date_value is None:
                continue
            previous_description = fold_text(ws.cell(previous, columns["description"]).value)
            credit_value = ws.cell(previous, columns["credit"]).value
            if "THANH TOAN" in previous_description and credit_value not in (None, "", 0, 0.0):
                transaction_row = previous
            break

        if transaction_row is None:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "KIOT_PAYMENT_ABOVE_KEYWORD_NOT_FOUND",
                    f"Có từ khóa “{description}” nhưng không tìm thấy dòng Thanh toán hợp lệ ngay phía trên",
                    source,
                    ws.title,
                    row_number,
                )
            )
            continue
        transaction_for_keyword[transaction_row].append(description)

    for transaction_row, descriptions in sorted(transaction_for_keyword.items()):
        raw_datetime = ws.cell(transaction_row, columns["datetime"]).value
        payment_date = parse_date(raw_datetime)
        payment_time = parse_time(raw_datetime)
        if payment_date is None:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "KIOT_INVALID_PAYMENT_DATE",
                    "Không đọc được ngày của dòng Thanh toán",
                    source,
                    ws.title,
                    transaction_row,
                )
            )
            continue
        try:
            amount = money_to_int(ws.cell(transaction_row, columns["credit"]).value)
        except ValueError as exc:
            result.issues.append(
                _issue(
                    Severity.ERROR,
                    "KIOT_INVALID_PAYMENT_AMOUNT",
                    str(exc),
                    source,
                    ws.title,
                    transaction_row,
                )
            )
            continue
        unique_descriptions = list(dict.fromkeys(descriptions))
        if len(descriptions) > 1:
            result.issues.append(
                _issue(
                    Severity.WARNING,
                    "MULTIPLE_KEYWORD_LINES_FOR_PAYMENT",
                    "Nhiều dòng chứa từ khóa cùng liên kết với một dòng Thanh toán; nội dung đã được gộp",
                    source,
                    ws.title,
                    transaction_row,
                )
            )
        result.payments.append(
            PaymentRecord(
                source="KIOT",
                payment_date=payment_date,
                payment_time=payment_time,
                amount=amount,
                content=" | ".join(unique_descriptions),
                source_file=source.name,
                sheet=ws.title,
                source_row=transaction_row,
            )
        )
    if not result.payments and not any(issue.is_error for issue in result.issues):
        result.issues.append(
            _issue(
                Severity.WARNING,
                "KIOT_KEYWORD_NOT_FOUND",
                f"Không tìm thấy giao dịch có từ khóa “{result.keyword}” trong cột Diễn giải",
                source,
                ws.title,
            )
        )
    return result


def expected_months(period: ReportPeriod) -> set[tuple[int, int]]:
    if not period.start or not period.end:
        return set()
    months: set[tuple[int, int]] = set()
    year, month = period.start.year, period.start.month
    while (year, month) <= (period.end.year, period.end.month):
        months.add((year, month))
        month += 1
        if month == 13:
            year += 1
            month = 1
    return months


def hddt_covered_months(data: HddtData) -> set[tuple[int, int]]:
    months: set[tuple[int, int]] = set()
    for period in data.periods:
        months.update(expected_months(period))
    return months


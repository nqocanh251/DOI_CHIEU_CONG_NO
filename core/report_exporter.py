from __future__ import annotations

from collections import Counter
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import (
    HddtData,
    InvoiceReconciliation,
    KiotData,
    PaymentReconciliation,
    SmartData,
    ValidationIssue,
)
from core.normalization import format_date, format_time, is_advisory_invoice_status


NAVY = "173B57"
BLUE = "1F5F87"
BLUE_LIGHT = "EAF3F8"
SLATE = "526577"
SLATE_LIGHT = "F4F7FA"
GRID = "D9E3EB"
WHITE = "FFFFFF"
GREEN = "E7F4EC"
GREEN_TEXT = "1E694B"
YELLOW = "FFF4D6"
YELLOW_TEXT = "805400"
RED = "FDE9E7"
RED_TEXT = "9C342F"
SMART_MISSING = "DDEBF7"
SMART_MISSING_TEXT = "1F4E78"
KIOT_MISSING = "F4CCCC"
MONEY_FORMAT = '#,##0 "đ";[Red]-#,##0 "đ"'
REPORT_TITLE_FONT_SIZE = 18
REPORT_SECTION_FONT_SIZE = 12
REPORT_BODY_FONT_SIZE = 11
REPORT_DETAIL_ROW_HEIGHT = 42
REPORT_HEADER_ROW_HEIGHT = 40


def invoice_results_to_dataframe(reconciliation: InvoiceReconciliation | None) -> pd.DataFrame:
    columns = [
        "STT",
        "MST nhà cung cấp",
        "Tên nhà cung cấp tham khảo",
        "Ký hiệu hóa đơn",
        "Số hóa đơn",
        "Ngày HĐĐT",
        "Ngày Smart",
        "Tổng tiền HĐĐT",
        "Tổng tiền Smart",
        "Chênh lệch",
        "Trạng thái HĐĐT / cảnh báo",
        "Trạng thái",
        "Nguyên nhân",
        "File HĐĐT",
        "Dòng HĐĐT",
        "File Smart",
        "Dòng Smart",
    ]
    if reconciliation is None:
        return pd.DataFrame(columns=columns)
    rows = []
    for index, result in enumerate(reconciliation.results, 1):
        hddt = result.hddt
        smart = result.smart
        supplier = hddt or smart
        hddt_amount = hddt.amount if hddt else None
        smart_amount = smart.amount if smart else None
        rows.append(
            {
                "STT": index,
                "MST nhà cung cấp": supplier.supplier_tax_id if supplier else "",
                "Tên nhà cung cấp tham khảo": supplier.supplier_name if supplier else "",
                "Ký hiệu hóa đơn": hddt.invoice_symbol if hddt else "",
                "Số hóa đơn": supplier.invoice_number_raw if supplier else "",
                "Ngày HĐĐT": format_date(hddt.invoice_date) if hddt else "",
                "Ngày Smart": format_date(smart.invoice_date) if smart else "",
                "Tổng tiền HĐĐT": hddt_amount,
                "Tổng tiền Smart": smart_amount,
                "Chênh lệch": (hddt_amount - smart_amount) if hddt and smart else None,
                "Trạng thái HĐĐT / cảnh báo": hddt.status if hddt else "",
                "Trạng thái": result.status,
                "Nguyên nhân": result.reason,
                "File HĐĐT": hddt.source_file if hddt else "",
                "Dòng HĐĐT": hddt.source_row if hddt else None,
                "File Smart": smart.source_file if smart else "",
                "Dòng Smart": smart.source_row if smart else None,
            }
        )
    return pd.DataFrame(rows, columns=columns)


def payment_results_to_dataframe(reconciliation: PaymentReconciliation | None) -> pd.DataFrame:
    columns = [
        "STT",
        "Trạng thái",
        "Ngày Smart",
        "Chứng từ Smart",
        "Nội dung Smart",
        "Số tiền Smart",
        "Ngày KIOT",
        "Giờ KIOT",
        "Nội dung KIOT",
        "Số tiền KIOT",
        "Số ngày chênh lệch",
        "Nguyên nhân",
        "File Smart",
        "Dòng Smart",
        "File KIOT",
        "Dòng KIOT",
    ]
    if reconciliation is None:
        return pd.DataFrame(columns=columns)
    rows = []
    for index, result in enumerate(reconciliation.results, 1):
        smart = result.smart
        kiot = result.kiot
        rows.append(
            {
                "STT": index,
                "Trạng thái": result.status,
                "Ngày Smart": format_date(smart.payment_date) if smart else "",
                "Chứng từ Smart": smart.reference if smart else "",
                "Nội dung Smart": smart.content if smart else "",
                "Số tiền Smart": smart.amount if smart else None,
                "Ngày KIOT": format_date(kiot.payment_date) if kiot else "",
                "Giờ KIOT": format_time(kiot.payment_time) if kiot else "",
                "Nội dung KIOT": kiot.content if kiot else "",
                "Số tiền KIOT": kiot.amount if kiot else None,
                "Số ngày chênh lệch": result.day_difference,
                "Nguyên nhân": result.reason,
                "File Smart": smart.source_file if smart else "",
                "Dòng Smart": smart.source_row if smart else None,
                "File KIOT": kiot.source_file if kiot else "",
                "Dòng KIOT": kiot.source_row if kiot else None,
            }
        )
    dataframe = pd.DataFrame(rows, columns=columns)
    dataframe["Số ngày chênh lệch"] = pd.array(
        dataframe["Số ngày chênh lệch"], dtype="Int64"
    )
    return dataframe


def issues_to_dataframe(issues: list[ValidationIssue]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Mức độ": issue.severity.value,
                "Mã lỗi": issue.code,
                "Thông báo": issue.message,
                "File": issue.source_file,
                "Sheet": issue.sheet,
                "Dòng": issue.row,
            }
            for issue in issues
        ],
        columns=["Mức độ", "Mã lỗi", "Thông báo", "File", "Sheet", "Dòng"],
    )


def _append_dataframe(ws, dataframe: pd.DataFrame):
    ws.append(list(dataframe.columns))
    for row in dataframe.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])


def _status_style(status: str) -> tuple[PatternFill, Font]:
    normalized = status.upper().strip()
    if normalized == "KHỚP":
        return PatternFill("solid", fgColor=GREEN), Font(
            name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=GREEN_TEXT, bold=True
        )
    if "SMART THIẾU" in normalized:
        return PatternFill("solid", fgColor=SMART_MISSING), Font(
            name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=SMART_MISSING_TEXT, bold=True
        )
    if "KIOT THIẾU" in normalized:
        return PatternFill("solid", fgColor=KIOT_MISSING), Font(
            name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=RED_TEXT, bold=True
        )
    if any(
        marker in normalized
        for marker in (
            "CẦN KIỂM TRA",
            "TRÙNG",
            "THAY THẾ",
            "ĐIỀU CHỈNH",
            "WARNING",
            "CẢNH BÁO",
        )
    ):
        return PatternFill("solid", fgColor=YELLOW), Font(
            name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=YELLOW_TEXT, bold=True
        )
    return PatternFill("solid", fgColor=RED), Font(
        name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=RED_TEXT, bold=True
    )


def _set_report_printing(ws, *, landscape: bool = True):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.color = SLATE
    ws.oddFooter.right.text = "Đối chiếu công nợ"
    ws.oddFooter.right.size = 9
    ws.oddFooter.right.color = SLATE


def _column_width(header: str, calculated_width: int) -> int:
    exact_widths = {
        "STT": 7,
        "MST nhà cung cấp": 17,
        "Tên nhà cung cấp tham khảo": 34,
        "Ký hiệu hóa đơn": 16,
        "Số hóa đơn": 14,
        "Ngày HĐĐT": 14,
        "Ngày Smart": 14,
        "Ngày ngân hàng": 14,
        "Ngày KIOT": 14,
        "Giờ KIOT": 11,
        "Tổng tiền HĐĐT": 19,
        "Tổng tiền Smart": 19,
        "Số tiền Smart": 19,
        "Số tiền KIOT": 19,
        "Số tiền ngân hàng": 19,
        "Chênh lệch": 17,
        "Số ngày chênh lệch": 18,
        "Trạng thái HĐĐT / cảnh báo": 27,
        "Trạng thái": 24,
        "Nguyên nhân": 44,
        "Chứng từ Smart": 20,
        "Nội dung Smart": 42,
        "Nội dung KIOT": 42,
        "Nội dung chuyển khoản": 48,
        "ID giao dịch ngân hàng": 24,
        "ID chứng từ Smart": 19,
        "Tên khách hàng Smart": 36,
        "TK Nợ Smart": 15,
        "Số ứng viên Smart": 18,
        "Số HĐ ứng viên": 34,
        "Khách hàng ứng viên": 45,
        "File HĐĐT": 30,
        "File Smart": 30,
        "File KIOT": 30,
        "Dòng HĐĐT": 13,
        "Dòng Smart": 13,
        "Dòng KIOT": 13,
        "Mức độ": 14,
        "Mã lỗi": 23,
        "Thông báo": 58,
        "File": 32,
        "Sheet": 16,
        "Dòng": 10,
    }
    return exact_widths.get(header, min(max(calculated_width, 11), 38))


def _style_sheet(ws, status_column_name: str | None = "Trạng thái"):
    """Định dạng các sheet chi tiết nhưng giữ hàng 1 là hàng tiêu đề dữ liệu."""
    thin_grid = Side(style="thin", color=GRID)
    header_fill = PatternFill("solid", fgColor=NAVY)
    alternate_fill = PatternFill("solid", fgColor=SLATE_LIGHT)
    money_headers = {
        "Tổng tiền HĐĐT",
        "Tổng tiền Smart",
        "Chênh lệch",
        "Số tiền Smart",
        "Số tiền KIOT",
        "Số tiền ngân hàng",
        "Tổng tiền",
    }
    centered_headers = {
        "STT",
        "Ngày HĐĐT",
        "Ngày Smart",
        "Ngày ngân hàng",
        "Ngày KIOT",
        "Giờ KIOT",
        "Dòng HĐĐT",
        "Dòng Smart",
        "Dòng KIOT",
        "Dòng",
        "Số ngày chênh lệch",
        "Mức độ",
    }

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))
    ws.row_dimensions[1].height = REPORT_HEADER_ROW_HEIGHT
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    headers = {cell.column: str(cell.value or "") for cell in ws[1]}
    status_column = next(
        (column for column, header in headers.items() if header == status_column_name),
        None,
    )
    hddt_warning_column = next(
        (
            column
            for column, header in headers.items()
            if header == "Trạng thái HĐĐT / cảnh báo"
        ),
        None,
    )
    money_columns = {column for column, header in headers.items() if header in money_headers}
    integer_columns = {
        column for column, header in headers.items() if header == "Số ngày chênh lệch"
    }
    centered_columns = {column for column, header in headers.items() if header in centered_headers}

    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = REPORT_DETAIL_ROW_HEIGHT
        for column in range(1, ws.max_column + 1):
            cell = ws.cell(row_number, column)
            if row_number % 2 == 0:
                cell.fill = alternate_fill
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color="263746")
            horizontal = "center" if column in centered_columns else "left"
            if column in money_columns:
                horizontal = "right"
                cell.number_format = MONEY_FORMAT
            elif column in integer_columns:
                horizontal = "center"
                cell.number_format = "0"
            cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_grid)
        if status_column:
            status_cell = ws.cell(row_number, status_column)
            status_cell.fill, status_cell.font = _status_style(str(status_cell.value or ""))
            status_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if hddt_warning_column:
            warning_cell = ws.cell(row_number, hddt_warning_column)
            warning_text = str(warning_cell.value or "").upper()
            if any(marker in warning_text for marker in ("THAY THẾ", "ĐIỀU CHỈNH")):
                warning_cell.fill = PatternFill("solid", fgColor=YELLOW)
                warning_cell.font = Font(
                    name="Segoe UI",
                    size=REPORT_BODY_FONT_SIZE,
                    color=YELLOW_TEXT,
                    bold=True,
                )

    for column, header in headers.items():
        values = [
            str(ws.cell(row, column).value or "")
            for row in range(1, min(ws.max_row, 300) + 1)
        ]
        calculated = max((len(value) for value in values), default=0) + 2
        ws.column_dimensions[get_column_letter(column)].width = _column_width(header, calculated)

    _set_report_printing(ws)
    ws.print_title_rows = "1:1"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    ws.sheet_properties.tabColor = BLUE


def _summary_sections(
    hddt: HddtData | None,
    smart: SmartData,
    kiot: KiotData | None,
    invoice_reconciliation: InvoiceReconciliation | None,
    payment_reconciliation: PaymentReconciliation | None,
):
    invoice_counts = Counter(
        result.status for result in (invoice_reconciliation.results if invoice_reconciliation else [])
    )
    payment_counts = Counter(
        result.status for result in (payment_reconciliation.results if payment_reconciliation else [])
    )
    relevant_hddt_total = (
        sum(record.amount for record in hddt.invoices if record.supplier_tax_id == smart.supplier_tax_id)
        if hddt and invoice_reconciliation
        else None
    )
    smart_invoice_total = sum(record.amount for record in smart.invoices) if invoice_reconciliation else None
    smart_payment_total = sum(record.amount for record in smart.payments)
    kiot_payment_total = sum(record.amount for record in kiot.payments) if kiot else 0
    if invoice_reconciliation is not None and relevant_hddt_total is not None and smart_invoice_total is not None:
        invoice_rows = [
            ("Tổng tiền hóa đơn HĐĐT", relevant_hddt_total),
            ("Tổng tiền hóa đơn Smart", smart_invoice_total),
            ("Chênh lệch tổng hóa đơn", relevant_hddt_total - smart_invoice_total),
            (
                "Kết quả HĐĐT–Smart",
                "KHỚP HOÀN TOÀN"
                if invoice_reconciliation.is_fully_matched
                else "CẦN KIỂM TRA",
            ),
        ]
    else:
        invoice_rows = [
            ("Tổng tiền hóa đơn HĐĐT", "Chưa thực hiện"),
            ("Tổng tiền hóa đơn Smart", "Chưa thực hiện"),
            ("Chênh lệch tổng hóa đơn", "Chưa thực hiện"),
            ("Kết quả HĐĐT–Smart", "CHƯA THỰC HIỆN"),
        ]
    for status, count in sorted(invoice_counts.items()):
        invoice_rows.append((f"HĐĐT–Smart: {status}", count))

    payment_rows = [
        ("Tổng Đã thanh toán Smart", smart_payment_total),
        (
            "Tổng Ghi có KIOT có từ khóa",
            kiot_payment_total if kiot else "Chưa thực hiện",
        ),
        (
            "Chênh lệch tổng thanh toán",
            smart_payment_total - kiot_payment_total if kiot else "Chưa thực hiện",
        ),
    ]
    if payment_reconciliation is not None:
        payment_statuses = [result.status for result in payment_reconciliation.results]
        if payment_statuses and all(status == "KHỚP" for status in payment_statuses):
            step2_result = "KHỚP"
        elif payment_statuses:
            step2_result = "CẦN KIỂM TRA"
        else:
            step2_result = "KHÔNG CÓ GIAO DỊCH"
        payment_rows.append(("Kết quả Smart–KIOT", step2_result))
        for status, count in sorted(payment_counts.items()):
            payment_rows.append((f"Smart–KIOT: {status}", count))
    else:
        payment_rows.append(("Kết quả Smart–KIOT", "CHƯA THỰC HIỆN"))

    warning_rows = []
    if hddt and invoice_reconciliation:
        advisory_invoices = sorted(
            (
                record
                for record in hddt.invoices
                if record.supplier_tax_id == smart.supplier_tax_id
                and is_advisory_invoice_status(record.status)
            ),
            key=lambda record: (
                record.invoice_date or datetime.min.date(),
                record.invoice_number,
                record.source_row,
            ),
        )
        warning_rows = [
            (
                f"{record.status} · Số HĐ: {record.invoice_number_raw or record.invoice_number}",
                record.amount,
            )
            for record in advisory_invoices
        ]

    sections = [
        (
            "THÔNG TIN ĐỐI CHIẾU",
            [
                ("Thời điểm xuất báo cáo", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
                ("MST nhà cung cấp", smart.supplier_tax_id),
                ("Tên nhà cung cấp Smart", smart.supplier_name),
                ("Tên nhà cung cấp KIOT (tham khảo)", kiot.supplier_name if kiot else ""),
                ("Kỳ Smart", smart.period.label),
                ("Kỳ KIOT", kiot.period.label if kiot else "Chưa thực hiện"),
            ],
        ),
        (
            "NGUỒN DỮ LIỆU",
            [
                ("File HĐĐT", "; ".join(hddt.files) if hddt else "Chưa thực hiện"),
                ("File Smart", smart.source_file),
                ("File KIOT", kiot.source_file if kiot else "Chưa thực hiện"),
            ],
        ),
        ("KẾT QUẢ HĐĐT ↔ SMART", invoice_rows),
    ]
    if warning_rows:
        sections.append(("CẢNH BÁO HÓA ĐƠN HĐĐT · SỐ HĐ VÀ SỐ TIỀN", warning_rows))
    sections.append(("KẾT QUẢ SMART ↔ KIOT", payment_rows))
    return sections


def _style_summary_value(cell, label: str, value):
    if isinstance(value, (int, float)) and not label.startswith(("HĐĐT–Smart:", "Smart–KIOT:")):
        cell.number_format = MONEY_FORMAT
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    if label.startswith(("HĐĐT–Smart:", "Smart–KIOT:")):
        cell.number_format = "0"
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if label.startswith("Kết quả "):
        normalized = str(value or "").upper()
        if normalized in {"KHỚP", "KHỚP HOÀN TOÀN"}:
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=GREEN_TEXT, bold=True)
        elif normalized in {"CHƯA THỰC HIỆN", "KHÔNG CÓ GIAO DỊCH"}:
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=YELLOW_TEXT, bold=True)
        else:
            cell.fill = PatternFill("solid", fgColor=RED)
            cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=RED_TEXT, bold=True)
    warning_status = label.split(" · Số HĐ:", 1)[0]
    if is_advisory_invoice_status(warning_status):
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=YELLOW_TEXT, bold=True)


def _build_summary_sheet(
    ws,
    hddt: HddtData | None,
    smart: SmartData,
    kiot: KiotData | None,
    invoice_reconciliation: InvoiceReconciliation | None,
    payment_reconciliation: PaymentReconciliation | None,
):
    thin_grid = Side(style="thin", color=GRID)
    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = "BÁO CÁO ĐỐI CHIẾU CÔNG NỢ NHÀ CUNG CẤP"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Segoe UI", size=REPORT_TITLE_FONT_SIZE, color=WHITE, bold=True)
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells("A2:F2")
    subtitle = ws["A2"]
    supplier = smart.supplier_name or "Chưa xác định nhà cung cấp"
    subtitle.value = f"{supplier}  ·  MST: {smart.supplier_tax_id or 'Chưa xác định'}  ·  Kỳ: {smart.period.label}"
    subtitle.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    subtitle.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=NAVY, bold=True)
    subtitle.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 29
    ws.row_dimensions[3].height = 9

    current_row = 4
    for section_title, rows in _summary_sections(
        hddt,
        smart,
        kiot,
        invoice_reconciliation,
        payment_reconciliation,
    ):
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        section_cell = ws.cell(current_row, 1, section_title)
        section_cell.fill = PatternFill("solid", fgColor=BLUE)
        section_cell.font = Font(name="Segoe UI", size=REPORT_SECTION_FONT_SIZE, color=WHITE, bold=True)
        section_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        for offset, (label, value) in enumerate(rows):
            row = current_row
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            label_cell = ws.cell(row, 1, label)
            value_cell = ws.cell(row, 2, value)
            row_fill = PatternFill("solid", fgColor=SLATE_LIGHT if offset % 2 else WHITE)
            label_cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
            label_cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color=NAVY, bold=True)
            label_cell.alignment = Alignment(vertical="center", wrap_text=True)
            value_cell.fill = row_fill
            value_cell.font = Font(name="Segoe UI", size=REPORT_BODY_FONT_SIZE, color="263746")
            value_cell.alignment = Alignment(vertical="center", wrap_text=True)
            for column in range(1, 7):
                cell = ws.cell(row, column)
                cell.border = Border(bottom=thin_grid)
                if column >= 2 and cell.coordinate not in ws.merged_cells:
                    cell.fill = row_fill
            _style_summary_value(value_cell, label, value)
            ws.row_dimensions[row].height = 38 if label.startswith("File ") else 32
            current_row += 1
        current_row += 1

    ws.column_dimensions["A"].width = 39
    ws.column_dimensions["B"].width = 22
    for column in "CDEF":
        ws.column_dimensions[column].width = 14
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = NAVY
    ws.print_area = f"A1:F{current_row - 1}"
    _set_report_printing(ws, landscape=False)


def build_excel_report(
    hddt: HddtData | None,
    smart: SmartData,
    invoice_reconciliation: InvoiceReconciliation | None,
    kiot: KiotData | None = None,
    payment_reconciliation: PaymentReconciliation | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.properties.title = "Báo cáo đối chiếu công nợ nhà cung cấp"
    workbook.properties.subject = "Đối chiếu HĐĐT – Smart Pro – KIOT Việt"
    workbook.properties.creator = "Tool đối chiếu công nợ"
    workbook.properties.description = "Báo cáo được tạo tự động; dữ liệu nguồn không bị thay đổi."
    summary = workbook.active
    summary.title = "Tong_quan"
    _build_summary_sheet(
        summary,
        hddt,
        smart,
        kiot,
        invoice_reconciliation,
        payment_reconciliation,
    )

    invoice_sheet = workbook.create_sheet("HDDT_Smart")
    _append_dataframe(invoice_sheet, invoice_results_to_dataframe(invoice_reconciliation))
    _style_sheet(invoice_sheet)
    invoice_sheet.sheet_properties.tabColor = BLUE

    payment_sheet = workbook.create_sheet("Smart_KIOT")
    _append_dataframe(payment_sheet, payment_results_to_dataframe(payment_reconciliation))
    _style_sheet(payment_sheet)
    payment_sheet.sheet_properties.tabColor = GREEN_TEXT

    all_issues = list(invoice_reconciliation.issues) if invoice_reconciliation else []
    if payment_reconciliation is not None:
        seen = {(i.severity, i.code, i.message, i.source_file, i.sheet, i.row) for i in all_issues}
        for issue in payment_reconciliation.issues:
            key = (issue.severity, issue.code, issue.message, issue.source_file, issue.sheet, issue.row)
            if key not in seen:
                all_issues.append(issue)
                seen.add(key)
    error_sheet = workbook.create_sheet("Loi_du_lieu")
    _append_dataframe(error_sheet, issues_to_dataframe(all_issues))
    _style_sheet(error_sheet, status_column_name="Mức độ")
    error_sheet.sheet_properties.tabColor = RED_TEXT

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
